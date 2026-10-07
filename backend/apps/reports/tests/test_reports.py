"""Rapports et fil d'activité (plan L8, N13 et N14)."""

from __future__ import annotations

import datetime as dt
import io
from decimal import Decimal

import pytest
from django.utils import timezone

from apps.accounts.roles import OcFunction, Role
from apps.accounts.tests.factories import VerifiedUserFactory
from apps.accounts.tests.roles_helpers import client_for, make_member
from apps.conferences.tests.factories import EditionFactory
from apps.core.actor import Actor
from apps.core.models import AuditLog
from apps.reports import services
from apps.reports.activity import feed

pytestmark = pytest.mark.django_db

COMMAND = Actor.command("cli:test")


@pytest.fixture
def edition():
    today = timezone.localdate()
    return EditionFactory(
        timezone="Africa/Abidjan",
        start_date=today - dt.timedelta(days=1),
        end_date=today + dt.timedelta(days=1),
    )


def table(tables, key):
    return next(item for item in tables if item.key == key)


def test_n13_submission_report_counts_rates_without_names(edition):
    from apps.reviews.tests.helpers import in_status
    from apps.submissions.models import SubmissionAuthor, SubmissionStatus
    from apps.submissions.tests.factories import complete_submission

    accepted = in_status(complete_submission(edition), SubmissionStatus.ACCEPTED)
    rejected = in_status(complete_submission(edition), SubmissionStatus.REJECTED)
    complete_submission(edition)  # brouillon : hors des taux
    SubmissionAuthor.objects.filter(submission=accepted).update(country="CI")
    SubmissionAuthor.objects.filter(submission=rejected).update(country="SN")
    tables = services.submissions_section(edition)
    by_country = table(tables, "by_country")
    assert by_country.rows == [["CI", 1, 1, Decimal("100.0")], ["SN", 1, 0, Decimal("0.0")]]
    assert sum(row[1] for row in table(tables, "by_track").rows) == 2
    names = {accepted.submitter.profile.last_name, rejected.submitter.email}
    assert not any(name in str(tables) for name in names)


def test_n13_registration_finance_and_attendance_reports(edition):
    from apps.events.tests.certificate_helpers import present
    from apps.registrations.models import RegistrationStatus
    from apps.registrations.tests.factories import make_registration

    registration = present(edition)
    make_registration(edition, VerifiedUserFactory(), status=RegistrationStatus.PENDING)
    statuses = dict(table(services.registrations_section(edition), "by_status").rows)
    assert statuses["confirmée"] == 1 and statuses["en attente de paiement"] == 1
    by_day = table(services.attendance_section(edition), "by_day").rows
    assert by_day == [[timezone.localdate().isoformat(), 1, 1, Decimal("100.0")]]
    summary = table(services.finance_section(edition), "summary")
    assert summary.columns[0] == "Indicateur"
    assert registration.user.email not in str(services.registrations_section(edition))


def test_n13_satisfaction_report_respects_the_threshold(edition):
    from apps.surveys.tests.test_surveys import full_answers, open_survey

    survey, people = open_survey(edition, people=5)
    from apps.surveys import services as surveys

    for person in people[:4]:
        surveys.answer(survey, person, full_answers(survey, note=4))
    tables = services.satisfaction_section(edition)
    assert table(tables, "rates").rows == [["Votre avis", 5, 4, Decimal("80.0")]]
    assert table(tables, "ratings").rows == []  # moins de 5 réponses
    surveys.answer(survey, people[4], full_answers(survey, note=2))
    ratings = table(services.satisfaction_section(edition), "ratings").rows
    assert ratings[0][2:] == [5, Decimal("3.60")]


def test_n13_sections_follow_existing_capabilities_and_exports_are_journaled(edition):
    finance = make_member(edition, Role.OC_MEMBER, oc_function=OcFunction.FINANCE)
    client = client_for(finance)
    base = f"/v1/manage/editions/{edition.pk}/reports"
    codes = [row["code"] for row in client.get(base).json()]
    # Le CO « finances » lit les soumissions (F10), pas la relecture ni la satisfaction.
    assert codes == ["submissions", "registrations", "finance", "attendance", "budget", "sponsors"]
    assert client.get(f"{base}/reviews").status_code == 403
    assert client.get(f"{base}/satisfaction").status_code == 403
    assert client.get(f"{base}/inconnue").status_code == 404
    body = client.get(f"{base}/budget").json()
    assert body["code"] == "budget" and body["tables"][0]["key"] == "totals"
    for file_format, kind in (
        ("csv", "text/csv"),
        ("xlsx", "spreadsheetml"),
        ("pdf", "application/pdf"),
    ):
        response = client.get(f"{base}/budget/export?file_format={file_format}")
        assert response.status_code == 200 and kind in response["Content-Type"]
    pdf = client.get(f"{base}/budget/export?file_format=pdf").content
    assert pdf.startswith(b"%PDF")
    xlsx = client.get(f"{base}/budget/export?file_format=xlsx").content
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(xlsx))
    assert len(workbook.sheetnames) == 2  # une feuille par tableau
    entries = AuditLog.objects.filter(action="report.exported")
    assert entries.count() == 5 and entries.first().after["section"] == "budget"
    assert client.get(f"{base}/budget/export?file_format=doc").status_code == 400


def test_n14_activity_feed_is_a_whitelist_without_values(edition):
    from apps.logistics.services import tasks

    task = tasks.create_task(edition, {"title": "Réserver la salle"}, actor=COMMAND)
    from apps.core.audit import record

    record("dietary.exported", actor=COMMAND, edition=edition, after={"rows": 3})
    record("budget.line_updated", actor=COMMAND, edition=edition, after={"planned": "1000"})
    entries = feed(edition)
    assert [entry["action"] for entry in entries] == ["task.created"]
    assert set(entries[0]) == {"at", "action", "actor", "object_type", "object_id"}
    assert entries[0]["object_id"] == str(task.pk)
    member = make_member(edition, Role.OC_MEMBER, oc_function=OcFunction.PROGRAM)
    response = client_for(member).get(f"/v1/manage/editions/{edition.pk}/activity")
    assert response.status_code == 200 and len(response.json()) == 1
    reviewer = make_member(edition, Role.SC_MEMBER)
    assert client_for(reviewer).get(f"/v1/manage/editions/{edition.pk}/activity").status_code == 403


def test_n8_meal_order_exports_as_pdf_without_names(edition):
    from apps.logistics.services import meals

    meals.create_meal(edition, {"day": timezone.localdate(), "kind": "lunch"}, actor=COMMAND)
    logistics = make_member(edition, Role.OC_MEMBER, oc_function=OcFunction.LOGISTICS)
    response = client_for(logistics).get(
        f"/v1/manage/editions/{edition.pk}/logistics/meals/export?file_format=pdf"
    )
    assert response.status_code == 200 and response.content.startswith(b"%PDF")
    assert AuditLog.objects.get(action="meal.exported").after["format"] == "pdf"
