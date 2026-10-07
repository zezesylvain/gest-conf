"""Budget prévisionnel et réalisé (plan L8, N4) : service et API."""

from __future__ import annotations

import io
from decimal import Decimal

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.roles import OcFunction, Role
from apps.accounts.tests.roles_helpers import client_for, make_member
from apps.core.actor import Actor
from apps.core.errors import Invalid
from apps.core.models import AuditLog
from apps.logistics.models import BudgetLine, BudgetSource
from apps.logistics.services import budget as service
from apps.payments.services import manual
from apps.registrations.services import orders
from apps.registrations.tests.helpers import grid, open_edition, participant_user

pytestmark = pytest.mark.django_db

COMMAND = Actor.command("cli:test")
PDF = b"%PDF-1.7\n%%EOF\n"


@pytest.fixture
def edition():
    edition = open_edition()
    grid(edition, "etudiant", early_local="25000", early_international="50000")
    return edition


@pytest.fixture
def finance(edition):
    return make_member(edition, Role.OC_MEMBER, oc_function=OcFunction.FINANCE)


def paid(edition, amount="25000"):
    registration = orders.place_order(
        edition, participant_user(), category="etudiant", method="transfer", actor=COMMAND
    )
    manual.record_manual_payment(
        registration,
        method="transfer",
        amount=Decimal(amount),
        reference="VIR-1",
        received_on=timezone.localdate(),
        actor=COMMAND,
    )
    return registration


def test_n4_line_validation_amounts_exact_in_the_edition_currency(edition):
    """XOF sans décimale ; montants positifs ; poste de la bonne nature ; champs fermés."""
    cases = [
        ({"kind": "expense", "category": "venue", "label": "Salle", "planned": "10.5"}, "planned"),
        ({"kind": "expense", "category": "venue", "label": "Salle", "planned": "-1"}, "planned"),
        ({"kind": "expense", "category": "grants", "label": "Salle"}, "category"),
        ({"kind": "income", "category": "venue", "label": "Salle"}, "category"),
        ({"kind": "other", "category": "venue", "label": "Salle"}, "kind"),
        ({"kind": "expense", "category": "venue", "label": ""}, "label"),
        ({"kind": "expense", "category": "venue", "label": "x", "source": "sponsors"}, "source"),
    ]
    for data, field in cases:
        with pytest.raises(Invalid) as error:
            service.create_line(edition, data, actor=COMMAND)
        assert field in error.value.fields, data
    line = service.create_line(
        edition,
        {"kind": "expense", "category": "venue", "label": " Salle ", "planned": "500000"},
        actor=COMMAND,
    )
    assert (line.label, line.planned, line.source) == ("Salle", Decimal("500000"), "manual")


def test_n4_registrations_actual_is_computed_from_payments_and_read_only(edition):
    """La ligne « inscriptions » existe d'office ; son réalisé est l'encaissé net (L6) ; on
    n'y saisit ni le réalisé, ni la nature, ni le poste ; elle ne se supprime pas."""
    paid(edition)
    paid(edition)
    rows = service.lines(edition)
    computed = [line for line, _actual in rows if line.source == BudgetSource.REGISTRATIONS]
    assert len(computed) == 1
    line = computed[0]
    assert service.actual_of(line) == Decimal("50000")
    with pytest.raises(Invalid) as error:
        service.update_line(line, {"actual": "1", "kind": "expense"}, actor=COMMAND)
    assert set(error.value.fields) == {"actual", "kind"}
    service.update_line(line, {"planned": "2000000", "note": "200 inscrits"}, actor=COMMAND)
    with pytest.raises(Invalid):
        service.delete_line(line, actor=COMMAND)
    # Lecture répétée : toujours une seule ligne calculée.
    service.lines(edition)
    assert BudgetLine.objects.filter(edition=edition, source="registrations").count() == 1


def test_n4_summary_totals_and_balances(edition):
    paid(edition)
    service.create_line(
        edition,
        {"kind": "expense", "category": "venue", "label": "Salle", "planned": "300000"},
        actor=COMMAND,
    )
    service.create_line(
        edition,
        {
            "kind": "expense",
            "category": "catering",
            "label": "Pauses",
            "planned": "100000",
            "actual": "120000",
        },
        actor=COMMAND,
    )
    service.create_line(
        edition,
        {"kind": "income", "category": "grants", "label": "Ministère", "planned": "400000"},
        actor=COMMAND,
    )
    summary = service.summary(edition)
    assert summary["currency"] == "XOF"
    assert summary["expense"] == {"planned": Decimal("400000"), "actual": Decimal("120000")}
    assert summary["income"] == {"planned": Decimal("400000"), "actual": Decimal("25000")}
    assert summary["balance_planned"] == Decimal("0")
    assert summary["balance_actual"] == Decimal("-95000")
    # La ligne calculée des partenariats (L8.3) figure aussi, à zéro.
    assert [row["category"] for row in summary["by_category"]] == [
        "venue",
        "catering",
        "registrations",
        "sponsorship",
        "grants",
    ]


def test_n4_lines_are_audited(edition):
    line = service.create_line(
        edition,
        {"kind": "expense", "category": "venue", "label": "Salle", "planned": "300000"},
        actor=COMMAND,
    )
    service.update_line(line, {"actual": "310000"}, actor=COMMAND)
    service.delete_line(line, actor=COMMAND)
    actions = list(
        AuditLog.objects.filter(edition=edition, action__startswith="budget.")
        .order_by("id")
        .values_list("action", flat=True)
    )
    assert actions == ["budget.line_created", "budget.line_updated", "budget.line_deleted"]
    updated = AuditLog.objects.get(action="budget.line_updated")
    assert (updated.before, updated.after) == ({"actual": None}, {"actual": "310000"})


def test_n4_proof_is_a_private_file_checked_by_content(edition, settings, tmp_path):
    settings.GESTCONF_PRIVATE_FILES_DIR = str(tmp_path)
    line = service.create_line(
        edition, {"kind": "expense", "category": "venue", "label": "Salle"}, actor=COMMAND
    )
    with pytest.raises(Invalid):
        service.upload_proof(line, data=b"<html>", name="f.html", actor=COMMAND)
    line = service.upload_proof(line, data=PDF, name="facture.pdf", actor=COMMAND)
    assert (line.proof_kind, line.proof_name) == ("pdf", "facture.pdf")
    assert service.PROOFS.read(line.proof_storage_name) == PDF
    assert service.known_proofs() == [line.proof_storage_name]
    line = service.remove_proof(line, actor=COMMAND)
    assert line.proof_storage_name == "" and service.known_proofs() == []


def test_n4_budget_api_rights_and_export(edition, finance):
    """Budget : CO « finances » ; export CSV et XLSX journalisé, sous réauthentification."""
    paid(edition)
    base = f"/v1/manage/editions/{edition.pk}/budget"
    client = client_for(finance)
    response = client.post(
        f"{base}/lines",
        {"kind": "expense", "category": "venue", "label": "=SOMME(A1)", "planned": "300000"},
        format="json",
    )
    assert response.status_code == 201
    body = response.json()
    assert body["currency"] == "XOF"
    registrations = next(line for line in body["lines"] if line["source"] == "registrations")
    assert (registrations["computed"], registrations["actual"]) == (True, "25000.00")
    venue = next(line for line in body["lines"] if line["source"] == "manual")
    assert venue["variance"] is None and venue["proof"] is None
    response = client.get(f"{base}/export?file_format=xlsx")
    assert response.status_code == 200
    rows = list(load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True))
    assert rows[0][0] == "Nature"
    assert any(row[2] == "=SOMME(A1)" for row in rows[1:])
    response = client.get(f"{base}/export")
    assert response.status_code == 200
    assert "'=SOMME(A1)" in response.content.decode()
    assert AuditLog.objects.filter(action="budget.exported").count() == 2
    response = client_for(finance, recent_auth=False).get(f"{base}/export")
    assert (response.status_code, response.json()["code"]) == (403, "reauthentication_required")
    response = client.get(f"{base}/export?file_format=pdf")
    assert response.status_code == 400
    # Le CO « logistique » ne voit pas le budget.
    other = make_member(edition, Role.OC_MEMBER, oc_function=OcFunction.LOGISTICS)
    assert client_for(other).get(base).status_code == 403
