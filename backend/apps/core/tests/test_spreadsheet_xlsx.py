"""Exports XLSX (plan L8, N13 ; bilan de L8.0) : aucune chaîne ne devient une formule."""

from __future__ import annotations

import io
import re
import zipfile
from decimal import Decimal

from openpyxl import load_workbook

from apps.core.spreadsheet import XLSX_CONTENT_TYPE, xlsx_bytes, xlsx_response

DANGEROUS = ["=1+1", "+33 6 12", "-5", "@SUM(A1)", '=HYPERLINK("http://x","y")', "\tx"]


def test_n13_xlsx_strings_never_become_formulas():
    content = xlsx_bytes(
        ["Texte", "Montant"], [[text, Decimal("12500.50")] for text in DANGEROUS], title="Export"
    )
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        sheet = next(name for name in archive.namelist() if name.startswith("xl/worksheets/"))
        xml = archive.read(sheet).decode()
    assert not re.search(r"<f>", xml)
    rows = list(load_workbook(io.BytesIO(content)).active.iter_rows(values_only=True))
    assert rows[0] == ("Texte", "Montant")
    assert [row[0] for row in rows[1:]] == DANGEROUS
    # Les montants restent des nombres (sommables dans le tableur).
    assert rows[1][1] == 12500.5


def test_n13_xlsx_response_is_private_and_typed():
    response = xlsx_response(xlsx_bytes(["a"], [[None], [True], [3]], title=""), "budget.xlsx")
    assert response["Content-Type"] == XLSX_CONTENT_TYPE
    assert response["Cache-Control"] == "private, no-store"
    assert response["Content-Disposition"] == 'attachment; filename="budget.xlsx"'
